import csv
import io
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook


POWER_FACTORS = {"w": 0.001, "kw": 1.0, "mw": 1000.0}


def convert_power_to_kw(value: object, unit: str) -> float:
    key = unit.strip().lower()
    if key not in POWER_FACTORS:
        raise ValueError("load_unit must be W, kW or MW")
    number = float(value)
    if number < 0:
        raise ValueError("load must be non-negative")
    return number * POWER_FACTORS[key]


def normalize_timestamp(value: object, timezone_name: str) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    else:
        text = str(value).strip()
        if not text:
            raise ValueError("timestamp is empty")
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    try:
        local_zone = ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError as exc:
        raise ValueError(f"unknown timezone: {timezone_name}") from exc
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=local_zone)
    return parsed.astimezone(timezone.utc).replace(tzinfo=None)


def read_tabular_file(filename: str, content: bytes) -> tuple[list[str], list[dict[str, object]]]:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        columns = list(reader.fieldnames or [])
        return columns, [dict(row) for row in reader]
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = next(iterator)
        except StopIteration:
            return [], []
        columns = [str(value).strip() if value is not None else "" for value in header]
        rows = [dict(zip(columns, values)) for values in iterator if any(value is not None for value in values)]
        workbook.close()
        return columns, rows
    raise ValueError("only .csv, .xlsx and .xlsm files are supported")
