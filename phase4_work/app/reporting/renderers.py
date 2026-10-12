from __future__ import annotations
from abc import ABC,abstractmethod
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import json
from html import escape
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

SECTIONS=["报告基本信息","数据来源与局限","项目/企业/园区基本情况","基础能源与设备情况","碳排放核算结果","碳排放结构分析","碳达峰与碳中和目标","低碳措施及路径推荐","多场景仿真推演","模型验证摘要","碳市场情景分析","CCER 与绿证分析","投资及经济性分析","风险评估","可行性结论","实施建议"]

def _trace(snapshot,report_id):
    inp=snapshot.get("input_snapshot_json",{});data=snapshot.get("report_data_json",{});p=data.get("provenance",{});sv=inp.get("scenario_version",{});st=inp.get("simulation_task",{})
    return {"报告编号":f"CN-RPT-{report_id:06d}","report_id":report_id,"scenario_id":st.get("scenario_id") or sv.get("scenario_id"),"scenario_version_id":snapshot.get("scenario_version_id"),"simulation_task_id":snapshot.get("simulation_task_id"),"数据时间范围":f"{st.get('start_time') or st.get('year') or '未提供'} 至 {st.get('end_time') or st.get('year') or '未提供'}","核算模型版本":((data.get("carbon_accounting") or {}).get("model_version") or "快照记录版本"),"主要排放因子来源":st.get("electricity_factor_source") or "见核算结果快照","市场参数快照":inp.get("carbon_market"),"风险参数快照":inp.get("risk_parameters"),"平台版本":"V1.0.0-alpha.8","生成日期":datetime.now().strftime("%Y-%m-%d")}

def _chart_specs(data):
    accounting=data.get("carbon_accounting") or {};summary=accounting.get("summary") or {};batch=((data.get("scenario_analysis") or {}).get("comparison_json") or {}).get("items",[]);market=data.get("carbon_market") or [];risks=data.get("risk_assessment") or []
    names=[str(x.get("scenario_version_id","场景")) for x in batch]
    return [
      ("年度碳排放趋势图",list((summary.get("annual_emissions") or {}).keys()),list((summary.get("annual_emissions") or {}).values()),"line"),
      ("能源碳排构成图",list((summary.get("by_energy") or {}).keys()),list((summary.get("by_energy") or {}).values()),"pie"),
      ("Scope 1 / Scope 2 构成图",list((summary.get("by_scope") or {}).keys()),list((summary.get("by_scope") or {}).values()),"pie"),
      ("多场景碳排对比图",names,[x.get("cumulative_carbon_emission_t") for x in batch],"bar"),
      ("多场景累计减排对比图",names,[x.get("cumulative_carbon_reduction_t") for x in batch],"bar"),
      ("多场景经济性对比图",names,[x.get("total_electricity_cost") for x in batch],"bar"),
      ("可再生能源发电/自用情况图",names,[x.get("renewable_generation_kwh") for x in batch],"bar"),
      ("碳价趋势图",[str(x.get("scenario_version_id","情景")) for x in market],[(((x.get("input_snapshot") or {}).get("market") or {}).get("carbon_price_per_t")) for x in market],"line"),
      ("风险等级汇总图",[str(x.get("scenario_version_id","风险")) for x in risks],[{"LOW":1,"MEDIUM":2,"HIGH":3}.get(x.get("overall_level"),0) for x in risks],"bar"),
    ]
def _charts(data,out):
    plt.rcParams["font.sans-serif"]=["Microsoft YaHei","SimHei","DejaVu Sans"];plt.rcParams["axes.unicode_minus"]=False;files=[]
    for i,(title,labels,values,kind) in enumerate(_chart_specs(data),1):
        pairs=[(a,float(b)) for a,b in zip(labels,values) if b is not None];fig,ax=plt.subplots(figsize=(7.2,3.3),dpi=130)
        if pairs:
            labels=[p[0] for p in pairs];values=[p[1] for p in pairs]
            if kind=="pie" and sum(values)>0:ax.pie(values,labels=labels,autopct="%1.1f%%")
            elif kind=="line":ax.plot(labels,values,marker="o",color="#2E74B5")
            else:ax.bar(labels,values,color="#2E74B5")
        else:ax.text(.5,.5,"快照中无可用数据",ha="center",va="center",transform=ax.transAxes,color="#666666");ax.set_xticks([]);ax.set_yticks([])
        ax.set_title(title);fig.tight_layout();p=Path(out)/f"chart_{i}.png";fig.savefig(p,bbox_inches="tight");plt.close(fig);files.append((title,p))
    return files
def _flatten(value):return json.dumps(value,ensure_ascii=False,indent=2,default=str) if isinstance(value,(dict,list)) else str(value if value is not None else "无数据")

class ReportRenderer(ABC):
    format:str
    @abstractmethod
    def render(self,snapshot:dict,output_path:Path,report_id:int,title:str):...

class DocxReportRenderer(ReportRenderer):
    format="DOCX"
    def render(self,snapshot,output_path,report_id,title):
      from .formal_project_report import build_docx
      return build_docx(snapshot,output_path,report_id,title)
      """Legacy renderer retained below for compatibility reference."""
      data=snapshot["report_data_json"];trace=_trace(snapshot,report_id)
      with TemporaryDirectory() as td:
        charts=_charts(data,td);doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1);sec.header_distance=sec.footer_distance=Inches(.492)
        styles=doc.styles;normal=styles["Normal"];normal.font.name="Microsoft YaHei";normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.1
        for name,size,before,after,color in [("Title",30,0,8,"203748"),("Heading 1",16,16,8,"2E74B5"),("Heading 2",13,12,6,"2E74B5"),("Heading 3",12,8,4,"1F4D78")]:
            st=styles[name];st.font.name="Microsoft YaHei";st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color);st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
        header=sec.header.paragraphs[0];header.text="碳中和可行性与风险评估报告 | V1.0.0-alpha.8";header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run("第 ");fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");footer._p.append(fld);footer.add_run(" 页")
        for _ in range(5):doc.add_paragraph()
        p=doc.add_paragraph(style="Title");p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(title)
        p=doc.add_paragraph("正式版 · 碳中和可行性与风险评估");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for k in ("报告编号","生成日期","平台版本"):
            p=doc.add_paragraph(f"{k}：{trace[k]}");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break();doc.add_heading("目录",level=1);p=doc.add_paragraph();toc=OxmlElement("w:fldSimple");toc.set(qn("w:instr"),'TOC \\o "1-3" \\h \\z \\u');p._p.append(toc);doc.add_page_break()
        doc.add_heading("报告基本信息",level=1);table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.style="Table Grid";table.autofit=False;table.columns[0].width=Inches(1.875);table.columns[1].width=Inches(4.625);table.rows[0].cells[0].text="追溯字段";table.rows[0].cells[1].text="冻结值"
        for k,v in trace.items():cells=table.add_row().cells;cells[0].text=str(k);cells[1].text=_flatten(v)
        content={"数据来源与局限":data.get("data_provenance"),"项目/企业/园区基本情况":data.get("project") or snapshot["input_snapshot_json"].get("scenario_version"),"基础能源与设备情况":data.get("energy_and_devices") or snapshot["input_snapshot_json"].get("simulation_task"),"碳排放核算结果":data.get("carbon_accounting"),"碳排放结构分析":data.get("carbon_accounting"),"碳达峰与碳中和目标":snapshot["input_snapshot_json"].get("scenario_version"),"低碳措施及路径推荐":data.get("path_recommendations"),"多场景仿真推演":data.get("scenario_analysis"),"模型验证摘要":data.get("model_validation"),"碳市场情景分析":data.get("carbon_market"),"CCER 与绿证分析":snapshot["input_snapshot_json"].get("carbon_market"),"投资及经济性分析":data.get("scenario_analysis"),"风险评估":data.get("risk_assessment"),"可行性结论":data.get("conclusions"),"实施建议":data.get("implementation_recommendations") or "依据风险评估和场景比较结果分阶段实施，并定期滚动复核。"}
        chart_map={x[0]:x[1] for x in charts};section_charts={"碳排放核算结果":["年度碳排放趋势图"],"碳排放结构分析":["能源碳排构成图","Scope 1 / Scope 2 构成图"],"多场景仿真推演":["多场景碳排对比图","多场景累计减排对比图","可再生能源发电/自用情况图"],"碳市场情景分析":["碳价趋势图"],"投资及经济性分析":["多场景经济性对比图"],"风险评估":["风险等级汇总图"]}
        for heading in SECTIONS[1:]:
            doc.add_heading(heading,level=1);doc.add_paragraph(_flatten(content.get(heading)))
            for chart in section_charts.get(heading,[]):doc.add_paragraph(chart,style="Heading 2");doc.add_picture(str(chart_map[chart]),width=Inches(6.2))
        output_path.parent.mkdir(parents=True,exist_ok=True);doc.save(output_path)

class PdfReportRenderer(ReportRenderer):
    format="PDF"
    def render(self,snapshot,output_path,report_id,title):
      from .formal_project_report import build_pdf
      return build_pdf(snapshot,output_path,report_id,title)
      """Legacy renderer retained below for compatibility reference."""
      data=snapshot["report_data_json"];trace=_trace(snapshot,report_id);output_path.parent.mkdir(parents=True,exist_ok=True);pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"));styles=getSampleStyleSheet();body=ParagraphStyle("CN",parent=styles["BodyText"],fontName="STSong-Light",fontSize=9,leading=14);h1=ParagraphStyle("CNH1",parent=body,fontSize=16,textColor=colors.HexColor("#2E74B5"),spaceBefore=12,spaceAfter=8);cover=ParagraphStyle("Cover",parent=h1,fontSize=24,alignment=TA_CENTER,spaceAfter=20)
      with TemporaryDirectory() as td:
        charts=_charts(data,td);story=[Spacer(1,120),Paragraph(escape(title),cover),Paragraph(f"报告编号：{trace['报告编号']}",ParagraphStyle("center",parent=body,alignment=TA_CENTER)),Paragraph(f"生成日期：{trace['生成日期']}",ParagraphStyle("center2",parent=body,alignment=TA_CENTER)),PageBreak(),Paragraph("报告基本信息",h1),Table([[Paragraph(escape(str(k)),body),Paragraph(escape(_flatten(v)).replace("\n","<br/>"),body)] for k,v in trace.items()],colWidths=[120,340],style=TableStyle([("GRID",(0,0),(-1,-1),.3,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP"),("BACKGROUND",(0,0),(0,-1),colors.HexColor("#F2F4F7"))]))]
        section_data=[data.get("data_provenance"),snapshot["input_snapshot_json"].get("scenario_version"),snapshot["input_snapshot_json"].get("simulation_task"),data.get("carbon_accounting"),data.get("carbon_accounting"),snapshot["input_snapshot_json"].get("scenario_version"),data.get("path_recommendations"),data.get("scenario_analysis"),data.get("carbon_market"),snapshot["input_snapshot_json"].get("carbon_market"),data.get("scenario_analysis"),data.get("risk_assessment"),data.get("conclusions"),data.get("implementation_recommendations") or "依据冻结结果分阶段实施。"]
        for heading,value in zip(SECTIONS[1:],section_data):story.extend([Paragraph(heading,h1),Paragraph(escape(_flatten(value)).replace("\n","<br/>"),body)])
        story.append(PageBreak());story.append(Paragraph("报告图表",h1))
        for caption,path in charts:story.extend([Paragraph(caption,h1),Image(str(path),width=460,height=211)])
        SimpleDocTemplate(str(output_path),pagesize=A4,rightMargin=56,leftMargin=56,topMargin=56,bottomMargin=56,title=title).build(story)
